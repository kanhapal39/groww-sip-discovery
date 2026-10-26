"""Tests for the Golden Dataset review workbook. No network."""

from __future__ import annotations

from openpyxl import load_workbook

from pipeline.export import read_approved, write_review_workbook


def tag_row(record_id: str, status: str = "ok") -> dict:
    return {
        "record_id": record_id, "source_url": f"https://play.google.com/x?reviewId={record_id}", "status": status,
        "app_id": "com.a", "year": 2025, "relevance": "yes", "behavior": "stop_sip", "evidence_quote": "stopped my SIP",
        "quote_status": "exact", "confidence": 0.9, "is_hypothesis": False,
    }


def test_workbook_rows_start_pending_and_only_approved_are_promoted(tmp_path) -> None:
    path = tmp_path / "review.xlsx"
    count = write_review_workbook([tag_row("a"), tag_row("b"), tag_row("c", status="failed")], {"a": "I stopped my SIP"}, path)
    assert count == 2

    wb = load_workbook(path)
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert [r[header.index("review_status")] for r in ws.iter_rows(min_row=2, values_only=True)] == ["pending", "pending"]
    assert ws.cell(2, header.index("source_url") + 1).hyperlink.target.endswith("reviewId=a")
    assert read_approved(path) == []

    ws.cell(2, header.index("review_status") + 1).value = "approved"
    wb.save(path)
    approved = read_approved(path)
    assert [r["record_id"] for r in approved] == ["a"]
    assert approved[0]["text"] == "I stopped my SIP"


def test_golden_sample_spreads_across_source_and_behavior() -> None:
    from pipeline.export import golden_sample

    rows = []
    for i in range(100):
        rows.append({**tag_row(f"ps{i}"), "app_id": "com.a", "behavior": "unclear"})
    for i in range(3):
        rows.append({**tag_row(f"rd{i}"), "app_id": None, "source": "reddit", "behavior": "pause_sip"})
    rows.append({**tag_row("bad", status="failed"), "behavior": "stop_sip"})
    sample = golden_sample(rows, 10, seed=1)
    assert len(sample) == 10
    assert sum(r["stratum"] == "reddit | pause_sip" for r in sample) == 3  # rare cell keeps all its rows
    assert "bad" not in {r["record_id"] for r in sample}
    assert {r["stratum"] for r in sample} == {"com.a | unclear", "reddit | pause_sip"}


def test_reviewed_count_protects_work_in_progress(tmp_path) -> None:
    from pipeline.export import reviewed_count

    path = tmp_path / "s.xlsx"
    write_review_workbook([tag_row("a"), tag_row("b")], {}, path)
    assert reviewed_count(path) == 0
    wb = load_workbook(path)
    wb.active.cell(2, 1).value = "rejected"
    wb.save(path)
    assert reviewed_count(path) == 1


def test_add_provenance_joins_clean_records(tmp_path) -> None:
    import json

    import pytest

    from pipeline.export import add_provenance

    clean = tmp_path / "clean.jsonl"
    clean.write_text(json.dumps({"record_id": "a", "source": "reddit", "collected_at": "2026-10-26T00:00:00Z", "context": "Paused SIP?",
                                 "language": "en", "metadata": {"kind": "comment", "subreddit": "mutualfunds", "published_at": "2024-01-02"}}), encoding="utf-8")
    (row,) = add_provenance([{"record_id": "a", "source_url": "https://x"}], clean)
    assert (row["source"], row["collected_at"], row["published_at"], row["subreddit"]) == ("reddit", "2026-10-26T00:00:00Z", "2024-01-02", "mutualfunds")
    with pytest.raises(ValueError):
        add_provenance([{"record_id": "missing", "source_url": "https://x"}], clean)
