"""Tests for the labelled evaluation set tooling. No network; workbooks are written to tmp_path."""

from __future__ import annotations

import yaml
from openpyxl import load_workbook

from research import ground_truth as gt
from research.opportunities import load_definitions
from tests.test_research import golden

CODEBOOK = yaml.safe_load((gt.PROJECT_ROOT / "codebook.yaml").read_text(encoding="utf-8"))
DEFINITIONS = load_definitions()
THEMES = [d["id"] for d in DEFINITIONS["opportunities"]]


def rows(n: int = 30) -> list[dict]:
    return [golden(f"id{i:02d}", f"text {i}", behavior="pause_sip", in_app_pain="mandate_autopay", unmet_need="easy_mandate_control")
            for i in range(n)]


def test_sample_is_reproducible_and_respects_exclusions() -> None:
    first = gt.sample(rows(), 20, seed=7, exclude=("id0",))
    assert [r["record_id"] for r in first] == [r["record_id"] for r in gt.sample(rows(), 20, seed=7, exclude=("id0",))]
    assert len(first) == 20 and not any(r["record_id"].startswith("id0") for r in first)
    assert [r["record_id"] for r in gt.sample(rows(), 20, seed=8)] != [r["record_id"] for r in first]


def test_sheet_hides_engine_tags_and_offers_codebook_values(tmp_path) -> None:
    path = tmp_path / "gt.xlsx"
    gt.write_sheet(path, rows(3), CODEBOOK, THEMES, {"seed": "7"})
    wb = load_workbook(path)
    header = [c.value for c in wb["labels"][1]]
    assert header == [*gt.SHOWN, *gt.LABELS]
    cells = {str(c.value) for row in wb["labels"].iter_rows(min_row=2) for c in row if c.value is not None}
    assert not cells & {"pause_sip", "mandate_autopay", "easy_mandate_control"}  # engine tags never shown
    assert {"tenure", "behavior"} <= {r[0] for r in wb["codebook"].iter_rows(values_only=True)}
    assert all(c.value in (None, "") for row in wb["labels"].iter_rows(min_row=2, min_col=len(gt.SHOWN) + 1) for c in row)


def test_compare_counts_agreement_per_field_with_own_denominator(tmp_path) -> None:
    data = rows(3)
    path = tmp_path / "gt.xlsx"
    gt.write_sheet(path, data, CODEBOOK, THEMES, {})
    wb = load_workbook(path)
    ws = wb["labels"]
    col = {c.value: c.column for c in ws[1]}
    ws.cell(2, col["behavior"], "pause_sip")                # agrees
    ws.cell(2, col["theme"], "mandate_reliability")         # engine matches this opportunity
    ws.cell(3, col["behavior"], "stop_sip")                 # disagrees
    ws.cell(3, col["tenure"], "new_under_1y")               # engine says unclear
    wb.save(path)
    result = gt.compare(gt.read_labels(path), data, DEFINITIONS)
    assert result["records_labelled"] == 2
    assert result["fields"]["behavior"] == {"labelled": 2, "agree": 1}
    assert result["fields"]["tenure"] == {"labelled": 1, "agree": 0}
    assert result["fields"]["theme"] == {"labelled": 1, "agree": 1}
    assert result["fields"]["unmet_need"] == {"labelled": 0, "agree": 0}
    assert {(d["record_id"], d["field"]) for d in result["disagreements"]} == {("id01", "behavior"), ("id01", "tenure")}
